"""Generate the real DOCX regression fixture for GitHub issue #2634."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

OUTPUT = Path(__file__).with_name("issue_2634_vertical_merge.docx")


def main() -> None:
    document = Document()
    document.add_heading("Genetic Screening Test Catalogue", level=1)
    document.add_paragraph(
        "Regression scenario: the method column is vertically merged, but "
        "every test row must still carry that method after conversion."
    )

    rows = [
        ("Q0101", "Hereditary Breast Cancer", "BRCA1", "15 working days"),
        ("Q0102", "Hereditary Ovarian Cancer", "BRCA2", "15 working days"),
        ("Q0103", "Lynch Syndrome", "MLH1", "15 working days"),
        ("Q0104", "Familial Adenomatous Polyposis", "APC", "15 working days"),
    ]
    method = (
        "Method: genomic DNA is extracted from peripheral blood and read by "
        "next-generation sequencing, with candidate variants confirmed by "
        "Sanger sequencing; results must be read alongside the clinical picture."
    )

    table = document.add_table(rows=len(rows) + 1, cols=5)
    table.style = "Table Grid"
    headers = ("Test Code", "Test Name", "Gene", "Turnaround", "Method")
    for column, value in enumerate(headers):
        table.cell(0, column).text = value

    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            table.cell(row_index, column).text = value

    merged_method = table.cell(1, 4).merge(table.cell(len(rows), 4))
    merged_method.text = method
    merged_method.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    document.add_paragraph(
        "Control table: adjacent cells may hold the same text, but they must "
        "not be treated as merged."
    )
    control = document.add_table(rows=2, cols=3)
    control.style = "Table Grid"
    for column, value in enumerate(("Column A", "Column B", "Column C")):
        control.cell(0, column).text = value
    for column, value in enumerate(("Same Value", "Same Value", "Distinct Value")):
        control.cell(1, column).text = value

    document.save(OUTPUT)
    print(f"generated={OUTPUT}")
    print(f"size={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
